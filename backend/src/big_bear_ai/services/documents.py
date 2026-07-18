from __future__ import annotations

import asyncio
import base64
import binascii
import io
import os
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from big_bear_ai.config import Settings
from big_bear_ai.database import Database
from big_bear_ai.repositories.resources import ReadOnlyResourceError


SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".json", ".csv", ".pdf", ".docx"})


class DocumentError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class DocumentService:
    def __init__(self, database: Database, settings: Settings) -> None:
        self.database = database
        self.settings = settings

    async def upload(self, payload: dict[str, Any]) -> dict[str, Any]:
        filename = payload.get("filename")
        if (
            not isinstance(filename, str)
            or not filename.strip()
            or Path(filename).name != filename
            or "/" in filename
            or "\\" in filename
        ):
            raise DocumentError("INVALID_FILE", "filename must be a plain file name")
        extension = Path(filename).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise DocumentError("UNSUPPORTED_FILE", f"Unsupported file type: {extension}")
        raw_content = payload.get("content_base64")
        if not isinstance(raw_content, str):
            raise DocumentError("INVALID_FILE", "content_base64 is required")
        try:
            content = base64.b64decode(raw_content, validate=True)
        except (binascii.Error, ValueError) as exc:
            raise DocumentError("INVALID_FILE", "content_base64 is invalid") from exc
        if not content:
            raise DocumentError("INVALID_FILE", "file must not be empty")
        if len(content) > self.settings.max_upload_bytes:
            raise DocumentError("FILE_TOO_LARGE", "file exceeds the configured upload limit")

        try:
            text = await asyncio.to_thread(_extract_text, extension, content)
        except Exception as exc:
            raise DocumentError("EXTRACTION_FAILED", "document text extraction failed") from exc
        text = _normalize_text(text)
        if not text:
            raise DocumentError("EXTRACTION_FAILED", "document does not contain extractable text")
        chunks = _chunk_text(text)

        identifier = str(uuid4())
        stored_name = f"{identifier}{extension}"
        await asyncio.to_thread(
            self.settings.uploads_dir.mkdir, parents=True, exist_ok=True
        )
        temporary_path = self.settings.uploads_dir / f".{stored_name}.tmp"
        final_path = self.settings.uploads_dir / stored_name
        await asyncio.to_thread(temporary_path.write_bytes, content)
        timestamp = _utc_now()

        def persist(connection: sqlite3.Connection) -> dict[str, Any]:
            connection.execute(
                """
                INSERT INTO documents(
                    id, title, description, filename, media_type, size_bytes,
                    extracted_text, index_status, file_path, author, read_only,
                    created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, 'ready', ?, '用户', 0, ?, ?)
                """,
                (
                    identifier,
                    str(payload.get("title") or Path(filename).stem),
                    str(payload.get("description") or ""),
                    filename,
                    str(payload.get("media_type") or _media_type(extension)),
                    len(content),
                    text,
                    str(final_path),
                    timestamp,
                    timestamp,
                ),
            )
            for ordinal, chunk in enumerate(chunks):
                chunk_id = str(uuid4())
                connection.execute(
                    "INSERT INTO document_chunks(id, document_id, ordinal, content) VALUES (?, ?, ?, ?)",
                    (chunk_id, identifier, ordinal, chunk),
                )
                connection.execute(
                    "INSERT INTO document_chunks_fts(chunk_id, document_id, content) VALUES (?, ?, ?)",
                    (chunk_id, identifier, chunk),
                )
            os.replace(temporary_path, final_path)
            return _get_document(connection, identifier)

        try:
            document = await self.database.run(persist)
        except Exception:
            for path in (temporary_path, final_path):
                path.unlink(missing_ok=True)
            raise
        document["chunk_count"] = len(chunks)
        return document

    async def get(self, document_id: str) -> dict[str, Any] | None:
        def operation(connection: sqlite3.Connection) -> dict[str, Any] | None:
            row = connection.execute(
                "SELECT * FROM documents WHERE id = ?", (document_id,)
            ).fetchone()
            return _decode_document(row) if row else None

        return await self.database.run(operation)

    async def list(
        self,
        *,
        search: str = "",
        limit: int = 50,
        cursor: str | None = None,
    ) -> dict[str, Any]:
        if not 1 <= limit <= 100:
            raise DocumentError("VALIDATION_ERROR", "limit must be between 1 and 100")
        try:
            offset = int(cursor or "0")
        except ValueError as exc:
            raise DocumentError("VALIDATION_ERROR", "cursor must be an integer offset") from exc
        if offset < 0:
            raise DocumentError("VALIDATION_ERROR", "cursor must not be negative")
        where = ""
        parameters: list[Any] = []
        if search:
            where = (
                " WHERE LOWER(title) LIKE ? OR LOWER(description) LIKE ? "
                "OR LOWER(COALESCE(filename, '')) LIKE ?"
            )
            parameters = [f"%{search.lower()}%"] * 3

        def operation(connection: sqlite3.Connection) -> dict[str, Any]:
            total = connection.execute(
                f"SELECT COUNT(*) FROM documents{where}", parameters
            ).fetchone()[0]
            rows = connection.execute(
                f"SELECT * FROM documents{where} "
                "ORDER BY read_only DESC, updated_at DESC, id LIMIT ? OFFSET ?",
                [*parameters, limit, offset],
            ).fetchall()
            next_offset = offset + len(rows)
            return {
                "items": [_decode_document(row) for row in rows],
                "total": total,
                "next_cursor": str(next_offset) if next_offset < total else None,
            }

        return await self.database.run(operation)

    async def search(
        self, query: str, *, limit: int = 8, document_ids: list[str] | None = None
    ) -> list[dict[str, Any]]:
        tokens = re.findall(r"\w+", query, flags=re.UNICODE)
        if not tokens:
            return []
        expression = " AND ".join(f'"{token.replace(chr(34), chr(34) * 2)}"' for token in tokens)

        def operation(connection: sqlite3.Connection) -> list[dict[str, Any]]:
            filters = ""
            identifiers: list[Any] = []
            if document_ids is not None:
                if not document_ids:
                    return []
                placeholders = ", ".join("?" for _ in document_ids)
                filters = f" AND c.document_id IN ({placeholders})"
                identifiers.extend(document_ids)
            exact_rows = connection.execute(
                """
                SELECT c.id AS chunk_id, c.document_id, c.content, d.title,
                       d.filename, 0.0 AS rank
                FROM document_chunks AS c
                JOIN documents AS d ON d.id = c.document_id
                WHERE LOWER(c.content) LIKE LOWER(?)
                """
                + filters
                + " ORDER BY c.ordinal LIMIT ?",
                [f"%{query}%", *identifiers, limit],
            ).fetchall()
            if exact_rows:
                return [dict(row) for row in exact_rows]

            fts_filters = filters.replace("c.document_id", "f.document_id")
            rows = connection.execute(
                """
                SELECT f.chunk_id, f.document_id, f.content, d.title, d.filename,
                       bm25(document_chunks_fts) AS rank
                FROM document_chunks_fts AS f
                JOIN documents AS d ON d.id = f.document_id
                WHERE document_chunks_fts MATCH ?
                """
                + fts_filters
                + " ORDER BY rank LIMIT ?",
                [expression, *identifiers, limit],
            ).fetchall()
            return [dict(row) for row in rows]

        return await self.database.run(operation)

    async def delete(self, document_id: str) -> None:
        def operation(connection: sqlite3.Connection) -> str | None:
            document = _get_document(connection, document_id)
            if document["read_only"]:
                raise ReadOnlyResourceError(f"document {document_id} is read-only")
            connection.execute(
                "DELETE FROM document_chunks_fts WHERE document_id = ?", (document_id,)
            )
            connection.execute("DELETE FROM documents WHERE id = ?", (document_id,))
            return document.get("file_path")

        file_path = await self.database.run(operation)
        if file_path:
            await asyncio.to_thread(Path(file_path).unlink, True)


def _get_document(connection: sqlite3.Connection, document_id: str) -> dict[str, Any]:
    row = connection.execute(
        "SELECT * FROM documents WHERE id = ?", (document_id,)
    ).fetchone()
    if row is None:
        raise KeyError(document_id)
    return _decode_document(row)


def _decode_document(row: sqlite3.Row) -> dict[str, Any]:
    document = dict(row)
    document["read_only"] = bool(document["read_only"])
    return document


def _extract_text(extension: str, content: bytes) -> str:
    if extension in {".txt", ".md", ".json", ".csv"}:
        return content.decode("utf-8")
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if extension == ".docx":
        document = Document(io.BytesIO(content))
        return "\n".join(_iter_docx_text(document))
    raise ValueError(extension)


def _iter_docx_text(container: Any):
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text:
                yield block.text
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    yield from _iter_docx_text(cell)


def _normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def _chunk_text(text: str, *, max_chars: int = 1200, overlap: int = 150) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks


def _media_type(extension: str) -> str:
    return {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".json": "application/json",
        ".csv": "text/csv",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }[extension]


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
